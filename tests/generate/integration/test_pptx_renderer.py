import logging
from pathlib import Path

import pymupdf
from pptx import Presentation
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_SHAPE_TYPE

from src.schemas import (
    InputAsset,
    LessonScript,
    PageBlock,
    PPTOutline,
    ScriptBlock,
    SlideOutline,
    VisualPlan,
)
from src.utils.renderers import pptx_renderer, render_ppt_outline


def _artifact_dir() -> Path:
    artifact_dir = Path(__file__).resolve().parent / "artifacts"
    artifact_dir.mkdir(parents=True, exist_ok=True)
    return artifact_dir


def _build_pdf_fixture(path: Path, page_count: int = 2) -> None:
    with pymupdf.open() as document:
        for page_number in range(page_count):
            page = document.new_page()
            page.insert_text((72, 72), f"Renderer source page {page_number + 1}", fontsize=20)
        document.save(path)


def _slide_text(slide) -> str:
    return "\n".join(shape.text for shape in slide.shapes if hasattr(shape, "text"))


def _count_auto_shapes(slide, auto_shape_type) -> int:
    count = 0
    for shape in slide.shapes:
        if shape.shape_type != MSO_SHAPE_TYPE.AUTO_SHAPE:
            continue
        if getattr(shape, "auto_shape_type", None) == auto_shape_type:
            count += 1
    return count


def test_render_ppt_outline_creates_pptx_file():
    artifact_dir = _artifact_dir()
    pdf_path = artifact_dir / "renderer_source.pdf"
    _build_pdf_fixture(pdf_path)

    ppt_outline = PPTOutline(
        deck_title="Renderer test deck",
        theme_name="default_teaching",
        slides=[
            SlideOutline(
                slide_id="slide_1",
                slide_type="cover",
                layout_template="cover_layout",
                title="Renderer test deck",
                subtitle="Intro",
                bullets=["Topic A", "Topic B"],
                speaker_notes="Cover notes",
                visual_plan=[VisualPlan(visual_type="none", source_unit_ids=[], caption="")],
                source_section_ids=["sec_1"],
                source_page_ids=["page_1"],
                source_unit_ids=["unit_1"],
                metadata={},
            ),
            SlideOutline(
                slide_id="slide_2",
                slide_type="concept",
                layout_template="content_layout",
                title="Core concept",
                subtitle=None,
                bullets=["Definition", "Scope", "Caution"],
                speaker_notes="Concept notes",
                visual_plan=[VisualPlan(visual_type="none", source_unit_ids=[], caption="")],
                source_section_ids=["sec_2"],
                source_page_ids=["page_2"],
                source_unit_ids=["unit_2"],
                metadata={},
            ),
        ],
        metadata={},
    )
    source_assets = [
        InputAsset(
            asset_id="asset_pdf",
            file_path=str(pdf_path),
            file_type="pdf",
            file_name=pdf_path.name,
        )
    ]
    source_pages = [
        PageBlock(
            page_id="page_2",
            course_id=None,
            lesson_id="LESSON_RENDER",
            section_id="sec_2",
            page=2,
            title="Core concept",
            content="Renderer source page 2",
            summary="Renderer source page 2",
            key_points=["Definition"],
            knowledge_points=["Definition"],
            source_unit_ids=["unit_2"],
            page_role="formula",
            file_id="asset_pdf",
            file_type="pdf",
        )
    ]
    lesson_script = LessonScript(
        lesson_title="Renderer test deck",
        script_blocks=[
            ScriptBlock(
                script_block_id="script_1",
                lesson_id="LESSON_RENDER",
                section_id="sec_1",
                title="Renderer test deck",
                page_range=[1],
                script_text="Cover script",
                key_points=["Topic A"],
                block_type="opening",
            )
        ],
        metadata={},
    )

    output_path = artifact_dir / "rendered_outline.pptx"
    rendered_path = render_ppt_outline(
        ppt_outline,
        output_path,
        lesson_script=lesson_script,
        source_assets=source_assets,
        source_pages=source_pages,
    )

    assert rendered_path.exists()
    presentation = Presentation(rendered_path)
    assert len(presentation.slides) == 2
    slide_texts = [_slide_text(slide) for slide in presentation.slides]
    assert any("Renderer test deck" in text for text in slide_texts)
    assert any(shape.shape_type == MSO_SHAPE_TYPE.PICTURE for shape in presentation.slides[1].shapes)


def test_render_specialized_layouts_add_distinct_shapes():
    artifact_dir = _artifact_dir()
    output_path = artifact_dir / "rendered_specialized_outline.pptx"

    ppt_outline = PPTOutline(
        deck_title="Specialized layouts",
        theme_name="default_teaching",
        slides=[
            SlideOutline(
                slide_id="slide_agenda",
                slide_type="agenda",
                layout_template="agenda_layout",
                title="Agenda",
                subtitle="Lesson roadmap",
                bullets=["Part 1", "Part 2", "Part 3", "Part 4"],
                speaker_notes="Roadmap notes",
                visual_plan=[VisualPlan(visual_type="none", source_unit_ids=[], caption="")],
                source_section_ids=["sec_1"],
                source_page_ids=[],
                source_unit_ids=[],
                metadata={},
            ),
            SlideOutline(
                slide_id="slide_process",
                slide_type="process",
                layout_template="process_layout",
                title="Process",
                subtitle="Execution flow",
                bullets=["Collect", "Clean", "Analyze", "Review"],
                speaker_notes="Process notes",
                visual_plan=[VisualPlan(visual_type="none", source_unit_ids=[], caption="")],
                source_section_ids=["sec_2"],
                source_page_ids=[],
                source_unit_ids=[],
                metadata={},
            ),
            SlideOutline(
                slide_id="slide_summary",
                slide_type="summary",
                layout_template="summary_layout",
                title="Summary",
                subtitle="Key takeaways",
                bullets=["Point A", "Point B", "Point C", "Point D", "Point E"],
                speaker_notes="Summary notes",
                visual_plan=[VisualPlan(visual_type="none", source_unit_ids=[], caption="")],
                source_section_ids=["sec_3"],
                source_page_ids=[],
                source_unit_ids=[],
                metadata={},
            ),
        ],
        metadata={},
    )

    rendered_path = render_ppt_outline(ppt_outline, output_path)
    presentation = Presentation(rendered_path)

    assert _count_auto_shapes(presentation.slides[0], MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE) >= 4
    assert _count_auto_shapes(presentation.slides[1], MSO_AUTO_SHAPE_TYPE.OVAL) >= 4
    assert _count_auto_shapes(presentation.slides[2], MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE) >= 2


def test_render_ppt_outline_surfaces_ppt_preview_warning(monkeypatch, caplog):
    artifact_dir = _artifact_dir()
    pptx_source_path = artifact_dir / "renderer_source_deck.pptx"
    Presentation().save(pptx_source_path)

    ppt_outline = PPTOutline(
        deck_title="Preview warning",
        theme_name="default_teaching",
        slides=[
            SlideOutline(
                slide_id="slide_example",
                slide_type="example",
                layout_template="example_layout",
                title="Example",
                subtitle="Source PPT preview",
                bullets=["Case context", "Key action"],
                speaker_notes="Preview notes",
                visual_plan=[VisualPlan(visual_type="none", source_unit_ids=[], caption="")],
                source_section_ids=["sec_1"],
                source_page_ids=["page_1"],
                source_unit_ids=[],
                metadata={},
            )
        ],
        metadata={},
    )
    source_assets = [
        InputAsset(
            asset_id="asset_pptx",
            file_path=str(pptx_source_path),
            file_type="pptx",
            file_name=pptx_source_path.name,
        )
    ]
    source_pages = [
        PageBlock(
            page_id="page_1",
            course_id=None,
            lesson_id="LESSON_RENDER",
            section_id="sec_1",
            page=1,
            title="Example",
            content="Example page",
            summary="Example page",
            key_points=["Case context"],
            knowledge_points=["Case context"],
            source_unit_ids=["asset_pptx_slide_1"],
            page_role="example",
            file_id="asset_pptx",
            file_type="pptx",
        )
    ]

    monkeypatch.setattr(pptx_renderer, "_find_soffice_path", lambda: None)

    output_path = artifact_dir / "rendered_warning_outline.pptx"
    with caplog.at_level(logging.WARNING):
        rendered_path = render_ppt_outline(
            ppt_outline,
            output_path,
            source_assets=source_assets,
            source_pages=source_pages,
        )

    presentation = Presentation(rendered_path)
    slide_text = _slide_text(presentation.slides[0])

    assert "LibreOffice not found" in caplog.text
    assert "LibreOffice unavailable for PPT preview" in slide_text
    assert not any(shape.shape_type == MSO_SHAPE_TYPE.PICTURE for shape in presentation.slides[0].shapes)
