import tempfile
from pathlib import Path

import pymupdf
from docx import Document
from pptx import Presentation


def build_sample_parser_assets(requested_dir: Path) -> dict[str, Path]:
    output_dir = _resolve_fixture_dir(requested_dir)

    txt_path = output_dir / "sample_generated_lesson.txt"
    pdf_path = output_dir / "sample_generated_lesson.pdf"
    pptx_path = output_dir / "sample_generated_lesson.pptx"
    docx_path = output_dir / "sample_generated_lesson.docx"

    _build_txt_fixture(txt_path)
    _build_pdf_fixture(pdf_path)
    _build_pptx_fixture(pptx_path)
    _build_docx_fixture(docx_path)

    return {
        "txt": txt_path,
        "pdf": pdf_path,
        "pptx": pptx_path,
        "docx": docx_path,
    }


def _build_txt_fixture(path: Path) -> None:
    path.write_text(
        "\n\n".join(
            [
                "平面假设\n平面假设认为梁在弯曲后，截面仍保持为平面。",
                "应力分布\n正应力沿截面高度线性变化，中性轴处应力为零。",
            ]
        ),
        encoding="utf-8",
    )


def _build_pdf_fixture(path: Path) -> None:
    document = pymupdf.open()

    first_page = document.new_page()
    first_page.insert_text(
        (72, 72),
        "材料力学导论\n课程介绍材料力学研究对象和核心问题。",
        fontsize=14,
    )

    second_page = document.new_page()
    second_page.insert_text(
        (72, 72),
        "弯曲正应力\n弯曲正应力与弯矩、截面惯性矩和距离中性轴的位置有关。",
        fontsize=14,
    )

    document.save(path)
    document.close()


def _build_pptx_fixture(path: Path) -> None:
    presentation = Presentation()

    title_slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    title_slide.shapes.title.text = "平面假设"
    title_slide.placeholders[1].text = "截面在变形后仍保持为平面。"

    content_slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    content_slide.shapes.title.text = "正应力公式"
    content_slide.placeholders[1].text = "sigma = M*y / I\n需要结合截面几何性质理解。"

    presentation.save(path)


def _build_docx_fixture(path: Path) -> None:
    document = Document()
    document.add_heading("剪力图", level=1)
    document.add_paragraph("剪力图反映梁内剪力沿长度方向的变化规律。")
    document.add_heading("弯矩图", level=1)
    document.add_paragraph("弯矩图可以帮助判断危险截面和最大正应力位置。")
    document.save(path)


def _resolve_fixture_dir(requested_dir: Path) -> Path:
    if requested_dir.exists() and requested_dir.is_dir():
        return requested_dir

    if not requested_dir.exists():
        try:
            requested_dir.mkdir(parents=True, exist_ok=True)
            return requested_dir
        except OSError:
            pass

    fallback_dir = (
        Path(tempfile.gettempdir()) / "ChaoxingAgentRuntime" / requested_dir.as_posix().replace(":", "")
    )
    fallback_dir.mkdir(parents=True, exist_ok=True)
    return fallback_dir
