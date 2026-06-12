from pptx import Presentation

from src.schemas import DocumentUnit, InputAsset
from src.utils.converters import convert_legacy_ppt_to_pptx

from .shared import infer_title, normalize_text


def extract_pptx_units(asset: InputAsset, logger=None) -> list[DocumentUnit]:
    presentation_path = asset.file_path
    converted_from_legacy_ppt = False
    if str(asset.file_path).lower().endswith(".ppt"):
        presentation_path = str(convert_legacy_ppt_to_pptx(asset.file_path, logger=logger))
        converted_from_legacy_ppt = True

    presentation = Presentation(presentation_path)
    units: list[DocumentUnit] = []

    for slide_number, slide in enumerate(presentation.slides, start=1):
        texts = []
        title = _extract_slide_title(slide)

        for shape in slide.shapes:
            shape_text = _extract_shape_text(shape)
            if shape_text:
                texts.append(shape_text)

        slide_text = "\n".join(texts).strip()
        if not slide_text:
            if logger is not None:
                logger.warning(
                    "Skipping empty pptx slide %d for asset %s",
                    slide_number,
                    asset.asset_id,
                )
            continue

        units.append(
            DocumentUnit(
                unit_id=f"{asset.asset_id}_slide_{slide_number}",
                unit_type="slide",
                index=slide_number,
                title=title or infer_title(slide_text, fallback=f"Slide {slide_number}"),
                text=slide_text,
                source_ref=f"slide {slide_number}",
                source_asset_id=asset.asset_id or "unknown_asset",
                metadata={
                    "file_name": asset.file_name,
                    "file_type": asset.file_type,
                    "source_index": slide_number,
                    "char_count": len(slide_text),
                    "converted_from_legacy_ppt": converted_from_legacy_ppt,
                    "presentation_path": presentation_path,
                },
            )
        )

    if logger is not None:
        logger.info("Parsed pptx asset %s into %d unit(s)", asset.asset_id, len(units))

    return units


def _extract_slide_title(slide) -> str:
    title_shape = getattr(slide.shapes, "title", None)
    if title_shape is None:
        return ""
    return _extract_shape_text(title_shape)


def _extract_shape_text(shape) -> str:
    if shape is None:
        return ""

    direct_text = getattr(shape, "text", None)
    if isinstance(direct_text, str):
        return normalize_text(direct_text)

    has_text_frame = getattr(shape, "has_text_frame", False)
    if has_text_frame and getattr(shape, "text_frame", None) is not None:
        parts = []
        for paragraph in shape.text_frame.paragraphs:
            runs_text = "".join(run.text for run in getattr(paragraph, "runs", []))
            paragraph_text = runs_text or getattr(paragraph, "text", "")
            paragraph_text = normalize_text(paragraph_text)
            if paragraph_text:
                parts.append(paragraph_text)
        return normalize_text("\n".join(parts))

    return ""
